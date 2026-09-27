"""Turn the assistant's Markdown into a Word file in the VPS Lakeshore 2.0 house style.

Mirrors the brand's document grammar (the `vps-*-docx` template set): DM Sans embedded in
the file, colour logo top-right, magenta small-caps kicker, navy title and rule, navy-header
tables with hairline rows, magenta clause numbers, cream panels, and a footer carrying the
document's status on the left and "n / N" on the right.
"""

from __future__ import annotations

import io
import os
import re
import uuid
import zipfile
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

# ------------------------------------------------------------------ brand tokens
NAVY, INK, MAGENTA, CREAM = "001E5F", "131944", "B5175A", "FFF8E1"
SLATE, RULE, PANEL, MAGENTA_TINT, WHITE = "5C6480", "D9DDE8", "F3F4F8", "F5E1EA", "FFFFFF"

FONT = {"regular": "DM Sans", "light": "DM Sans Light", "medium": "DM Sans Medium",
        "semibold": "DM Sans SemiBold", "italic": "DM Sans Italic"}
FONT_FILES = {"DM Sans": "DMSans.ttf", "DM Sans Light": "DMSansLight.ttf", "DM Sans Medium": "DMSansMedium.ttf",
              "DM Sans SemiBold": "DMSansSemiBold.ttf", "DM Sans Italic": "DMSansItalic.ttf"}

STATIC = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "static")
LOGO = os.path.join(STATIC, "brand", "logo-lakeshore-colour.png")
FONT_DIR = os.path.join(STATIC, "fonts")

CONTENT = 9866  # A4 with 1.8 cm side margins, in twips (DXA)

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[[^\]\n]{2,80}\](?!\())")
TABLE_SEP = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")

# Kicker, reference and footer for each kind of output.
KINDS = {
    "vet": ("LEGAL · DOCUMENT REVIEW", "Review note · Confidential · Check the points against the document before relying on it"),
    "answer": ("COMPANY SECRETARIAT · ADVICE NOTE", "Advice note · Check against the Act and the portal before acting"),
    "ask": ("COMPANY SECRETARIAT · ADVICE NOTE", "Advice note · Check against the Act and the portal before acting"),
    "draft": ("LEGAL · CONFIDENTIAL", "Template for guidance · Have Legal review before execution"),
}
GROUP_KICKER = {
    "Board & Shareholders": "GOVERNANCE · COMPANY SECRETARIAT",
    "Contracts": "LEGAL · CONFIDENTIAL",
    "Legal Letters & Disputes": "LEGAL · PRIVILEGED & CONFIDENTIAL",
    "Policies": "POLICY · DRAFT FOR APPROVAL",
}


# ------------------------------------------------------------------ low-level XML
# Word is strict about child order; these follow the ECMA-376 sequences.
ORDER = {
    "w:pPr": ["w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl",
              "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens",
              "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
              "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
              "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
              "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange"],
    "w:rPr": ["w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps", "w:strike",
              "w:dstrike", "w:outline", "w:shadow", "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid",
              "w:vanish", "w:webHidden", "w:color", "w:spacing", "w:w", "w:kern", "w:position", "w:sz", "w:szCs",
              "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs",
              "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath"],
    "w:tblPr": ["w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
                "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
                "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"],
}


def _reorder(el, kind: str) -> None:
    rank = {qn(t): i for i, t in enumerate(ORDER[kind])}
    kids = list(el)
    kids.sort(key=lambda c: rank.get(c.tag, len(rank)))
    for c in kids:
        el.remove(c)
    for c in kids:
        el.append(c)

def _el(tag: str, **attrs) -> OxmlElement:
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def _border(tag: str, color: str | None, size: int = 6):
    if color is None:
        return _el(tag, **{"w:val": "nil"})
    return _el(tag, **{"w:val": "single", "w:sz": size, "w:space": 0, "w:color": color})


def _cell_props(cell, width: int, fill: str | None = None, borders: dict | None = None,
                margins: tuple[int, int, int, int] = (0, 0, 0, 0), valign: str = "top"):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag in (qn("w:tcW"), qn("w:shd"), qn("w:tcBorders"), qn("w:tcMar"), qn("w:vAlign")):
            tcPr.remove(child)
    tcPr.append(_el("w:tcW", **{"w:w": width, "w:type": "dxa"}))
    b = _el("w:tcBorders")
    borders = borders or {}
    for side in ("top", "left", "bottom", "right"):
        spec = borders.get(side)
        b.append(_border(f"w:{side}", *(spec if spec else (None,))))
    tcPr.append(b)
    if fill:
        tcPr.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))
    mar = _el("w:tcMar")
    for side, val in zip(("top", "left", "bottom", "right"), margins):
        mar.append(_el(f"w:{side}", **{"w:w": val, "w:type": "dxa"}))
    tcPr.append(mar)
    tcPr.append(_el("w:vAlign", **{"w:val": valign}))


def _table(doc_or_cell, rows: int, widths: list[int], indent: int = 0):
    t = doc_or_cell.add_table(rows=rows, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = t._tbl.tblPr
    for child in list(tblPr):
        if child.tag in (qn("w:tblStyle"), qn("w:tblBorders"), qn("w:tblW"), qn("w:tblLayout")):
            tblPr.remove(child)
    tblPr.append(_el("w:tblW", **{"w:w": sum(widths), "w:type": "dxa"}))
    if indent:
        tblPr.append(_el("w:tblInd", **{"w:w": indent, "w:type": "dxa"}))
    bd = _el("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        bd.append(_border(f"w:{side}", None))
    tblPr.append(bd)
    tblPr.append(_el("w:tblLayout", **{"w:type": "fixed"}))
    mar = _el("w:tblCellMar")
    for side in ("left", "right"):
        mar.append(_el(f"w:{side}", **{"w:w": 0, "w:type": "dxa"}))
    tblPr.append(mar)
    _reorder(tblPr, "w:tblPr")
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    return t


def _para_border(p, side: str, color: str, size: int = 6, space: int = 1):
    pPr = p._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = _el("w:pBdr")
        pPr.append(bdr)
    bdr.append(_el(f"w:{side}", **{"w:val": "single", "w:sz": size, "w:space": space, "w:color": color}))
    _reorder(pPr, "w:pPr")


def _spacing(p, before: int = 0, after: int = 0, line: int = 300, keep_next: bool = False):
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Twips(before), Twips(after)
    pf.line_spacing = line / 240
    if keep_next:
        pf.keep_with_next = True
    _reorder(p._p.get_or_add_pPr(), "w:pPr")


# ------------------------------------------------------------------ runs
def _run(p, text: str, size: float = 10, weight: str = "regular", color: str = INK, caps: bool = False,
         tracking: int | None = None, italic: bool = False):
    # DM Sans has no rupee glyph; set it in Arial as the brand library does.
    for part in re.split(r"(₹)", text):
        if not part:
            continue
        r = p.add_run(part)
        name = "Arial" if part == "₹" else (FONT["italic"] if italic and weight == "regular" else FONT[weight])
        r.font.name = name
        rPr = r._r.get_or_add_rPr()
        fonts = rPr.find(qn("w:rFonts"))
        for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(k), name)
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor.from_string(color)
        if caps:
            r.font.all_caps = True
        if italic and weight != "regular":
            r.font.italic = True
        if tracking:
            rPr.append(_el("w:spacing", **{"w:val": tracking}))
        _reorder(rPr, "w:rPr")


def _inline(p, text: str, size: float = 10, color: str = INK, base: str = "regular"):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            _run(p, part[2:-2], size, "semibold", NAVY if color == INK else color)
        elif part.startswith("`") and part.endswith("`"):
            _run(p, part[1:-1], size, "medium", color)
        elif part.startswith("[") and part.endswith("]"):
            _run(p, part, size, "semibold", color)  # placeholders must be impossible to miss
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            _run(p, part[1:-1], size, base, color, italic=True)
        else:
            _run(p, part, size, base, color)


def _plain(text: str) -> str:
    return re.sub(r"\*\*|`", "", text).strip()


# ------------------------------------------------------------------ components
def _kicker(container, text: str):
    p = container.add_paragraph()
    _run(p, text, 7.5, "semibold", MAGENTA, caps=True, tracking=30)
    _spacing(p, after=80, line=240)


def _pill(container, text: str):
    w = max(1100, len(text) * 92 + 520)
    t = _table(container, 1, [w])
    c = t.cell(0, 0)
    line = (NAVY, 8)
    _cell_props(c, w, borders={"top": line, "bottom": line, "left": line, "right": line},
                margins=(70, 160, 70, 160), valign="center")
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, text, 8.5, "semibold", NAVY)
    _spacing(p, line=220)


def _masthead(doc, kicker: str, title: str, pill: str | None, meta: list[tuple[str, str]]):
    right = 2700
    t = _table(doc, 1, [CONTENT - right, right])
    left, logo = t.cell(0, 0), t.cell(0, 1)
    _cell_props(left, CONTENT - right, valign="bottom")
    _cell_props(logo, right)
    left.paragraphs[0]._p.getparent().remove(left.paragraphs[0]._p)
    _kicker(left, kicker)
    if pill:
        _pill(left, pill)
        _spacing(left.add_paragraph(), after=0, line=240)
    tp = left.add_paragraph()
    _run(tp, title, 20, "regular", NAVY)
    _spacing(tp, line=260)
    lp = logo.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if os.path.exists(LOGO):
        lp.add_run().add_picture(LOGO, width=Cm(4.3))
    _spacing(lp, line=240)
    rule = doc.add_paragraph()
    _para_border(rule, "bottom", NAVY, 6)
    _spacing(rule, before=100, after=200, line=120)
    if meta:
        w = CONTENT // len(meta)
        mt = _table(doc, 1, [w] * len(meta))
        for i, (lab, val) in enumerate(meta):
            c = mt.cell(0, i)
            _cell_props(c, w, margins=(0, 0, 140, 200 if i < len(meta) - 1 else 0))
            p = c.paragraphs[0]
            _run(p, lab, 7.5, "medium", SLATE, caps=True, tracking=30)
            _spacing(p, after=30, line=240)
            p2 = c.add_paragraph()
            _inline(p2, val, 10)
            _spacing(p2, line=260)
        _spacing(doc.add_paragraph(), after=80, line=240)


def _heading(doc, text: str, level: int):
    p = doc.add_paragraph(style="Heading 2" if level <= 2 else "Heading 3")
    if level <= 2:
        _run(p, _plain(text), 11.5, "semibold", NAVY)
        _spacing(p, before=220, after=80, line=260, keep_next=True)
    else:
        _run(p, _plain(text), 10.5, "medium", NAVY)
        _spacing(p, before=140, after=60, line=260, keep_next=True)


def _body(doc, text: str):
    p = doc.add_paragraph()
    _inline(p, text)
    _spacing(p, after=120)


def _bullet(doc, text: str, depth: int):
    p = doc.add_paragraph()
    left = 440 + depth * 360
    p.paragraph_format.left_indent = Twips(left)
    p.paragraph_format.first_line_indent = Twips(-260)
    p.paragraph_format.tab_stops.add_tab_stop(Twips(left))
    _run(p, "•\t", 10, "regular", NAVY)
    _inline(p, text)
    _spacing(p, after=60)


def _clause(doc, number: str, text: str, depth: int):
    """Numbered item: magenta number, a bold run-in title (if any) in navy semibold."""
    p = doc.add_paragraph()
    left = 560 + depth * 400
    p.paragraph_format.left_indent = Twips(left)
    p.paragraph_format.first_line_indent = Twips(-560)
    p.paragraph_format.tab_stops.add_tab_stop(Twips(left))
    _run(p, f"{number}.\t", 10, "semibold", MAGENTA)
    _inline(p, text)
    _spacing(p, after=80)


def _panel(doc, label: str | None, text: str, fill: str = CREAM, border: str | None = None, label_color: str = SLATE):
    t = _table(doc, 1, [CONTENT])
    c = t.cell(0, 0)
    side = (border, 8) if border else None
    _cell_props(c, CONTENT, fill=fill, margins=(160, 220, 160, 220),
                borders={"top": side, "bottom": side, "left": side, "right": side} if side else None)
    p = c.paragraphs[0]
    if label:
        _run(p, label, 7.5, "semibold", label_color, caps=True, tracking=30)
        _spacing(p, after=60, line=240)
        p = c.add_paragraph()
    _inline(p, text)
    _spacing(p, line=300)
    _spacing(doc.add_paragraph(), after=60, line=240)


def _data_table(doc, header: list[str], rows: list[list[str]]):
    n = len(header)
    first = 700 if n > 2 and all(re.fullmatch(r"\d{1,3}\.?", (r[0] if r else "").strip()) for r in rows if r) else None
    widths = [first] + [(CONTENT - first) // (n - 1)] * (n - 1) if first else [CONTENT // n] * n
    widths[-1] += CONTENT - sum(widths)
    t = _table(doc, 1 + len(rows), widths)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(_el("w:tblHeader"))
    for j, h in enumerate(header):
        c = t.cell(0, j)
        _cell_props(c, widths[j], fill=NAVY, margins=(150, 160, 150, 160), valign="center")
        _run(c.paragraphs[0], _plain(h), 9.5, "semibold", WHITE)
        _spacing(c.paragraphs[0], line=260)
    for i, r in enumerate(rows, start=1):
        for j in range(n):
            c = t.cell(i, j)
            _cell_props(c, widths[j], margins=(120, 160, 120, 160), valign="center", borders={"bottom": (RULE, 6)})
            val = r[j] if j < len(r) else ""
            p = c.paragraphs[0]
            if j == 0 and first:
                _run(p, _plain(val), 9.5, "semibold", MAGENTA)
            else:
                _inline(p, val, 9.5)
            _spacing(p, line=280)
    _spacing(doc.add_paragraph(), after=80, line=240)


def _cells(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def _footer(section, text: str):
    p = section.footer.paragraphs[0]
    p.paragraph_format.tab_stops.add_tab_stop(Twips(CONTENT), WD_TAB_ALIGNMENT.RIGHT)
    _para_border(p, "top", RULE, 6, 6)
    _run(p, text + "\t", 8.5, "regular", SLATE)
    for i, instr in enumerate(("PAGE", "NUMPAGES")):
        if i:
            _run(p, " / ", 8.5, "regular", SLATE)
        fld = _el("w:fldSimple", **{"w:instr": instr})
        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        rPr.append(_el("w:rFonts", **{"w:ascii": FONT["regular"], "w:hAnsi": FONT["regular"]}))
        rPr.append(_el("w:color", **{"w:val": SLATE}))
        rPr.append(_el("w:sz", **{"w:val": 17}))
        r.append(rPr)
        t = OxmlElement("w:t")
        t.text = "1"
        r.append(t)
        fld.append(r)
        p._p.append(fld)
    _spacing(p, line=300)


# ------------------------------------------------------------------ fonts
def _embed_fonts(data: bytes) -> bytes:
    """Embed DM Sans (obfuscated, as Word does) so the file looks right on PCs without the font."""
    files = {name: os.path.join(FONT_DIR, f) for name, f in FONT_FILES.items() if os.path.exists(os.path.join(FONT_DIR, f))}
    if not files:
        return data
    src = zipfile.ZipFile(io.BytesIO(data))
    out_buf = io.BytesIO()
    out = zipfile.ZipFile(out_buf, "w", zipfile.ZIP_DEFLATED)
    font_entries, rels = [], []
    for n, (name, path) in enumerate(files.items(), start=1):
        guid = str(uuid.uuid4()).upper()
        key = bytes.fromhex(guid.replace("-", ""))[::-1]
        raw = bytearray(open(path, "rb").read())
        for i in range(32):
            raw[i] ^= key[i % 16]
        out.writestr(f"word/fonts/font{n}.odttf", bytes(raw))
        rid = f"rIdLkFont{n}"
        rels.append(f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font" Target="fonts/font{n}.odttf"/>')
        font_entries.append(f'<w:font w:name="{name}"><w:charset w:val="00"/><w:family w:val="auto"/><w:pitch w:val="variable"/>'
                            f'<w:embedRegular r:id="{rid}" w:fontKey="{{{guid}}}"/></w:font>')
    for item in src.infolist():
        buf = src.read(item.filename)
        if item.filename == "word/fontTable.xml":
            buf = buf.decode().replace("</w:fonts>", "".join(font_entries) + "</w:fonts>").encode()
        elif item.filename == "[Content_Types].xml":
            buf = buf.decode().replace(
                "<Default Extension=\"xml\"",
                '<Default Extension="odttf" ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont"/><Default Extension="xml"', 1).encode()
        elif item.filename == "word/settings.xml":
            s = buf.decode()
            s = s.replace('<w:zoom w:val="bestFit"/>', '<w:zoom w:val="bestFit" w:percent="100"/>')
            s = re.sub(r"(<w:zoom[^>]*/>)", r"\1<w:embedTrueTypeFonts/>", s, count=1) if "<w:zoom" in s \
                else re.sub(r"(<w:settings[^>]*>)", r"\1<w:embedTrueTypeFonts/>", s, count=1)
            buf = s.encode()
        out.writestr(item, buf)
    out.writestr("word/_rels/fontTable.xml.rels",
                 '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                 '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">' + "".join(rels) + "</Relationships>")
    out.close()
    return out_buf.getvalue()


# ------------------------------------------------------------------ main
def reference_for(dept: str, code: str) -> str:
    return f"LHRC/{dept}/{code}/{date.today().year}/[NNN]"


def markdown_to_docx(markdown: str, title: str, company: str = "", kind: str = "draft",
                     group: str | None = None, ref: str | None = None) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Twips(11906), Twips(16838)
    sec.top_margin, sec.bottom_margin = Twips(1077), Twips(1250)
    sec.left_margin = sec.right_margin = Twips(1020)
    sec.header_distance, sec.footer_distance = Twips(500), Twips(560)

    normal = doc.styles["Normal"]
    normal.font.name = FONT["regular"]
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT["regular"])
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK)

    kicker, footer_text = KINDS.get(kind, KINDS["draft"])
    if kind == "draft" and group in GROUP_KICKER:
        kicker = GROUP_KICKER[group]

    lines = markdown.replace("\r\n", "\n").split("\n")
    # The first "# Title" becomes the masthead title.
    for idx, ln in enumerate(lines):
        if ln.strip():
            m = re.match(r"^#\s+(.*)$", ln.strip())
            if m:
                title = _plain(m.group(1))
                lines = lines[idx + 1:]
            break

    pill = None
    if ref:
        pill = f"{'Agreement no.' if group == 'Contracts' else 'Ref.'} {ref}"
    meta = [("Prepared by", "Company Secretariat & Legal"), ("Date", f"{date.today().day} {date.today():%B %Y}"),
            ("Status", "Draft for review" if kind == "draft" else "For internal use")]
    if company:
        meta.insert(0, ("Company", company))
    _masthead(doc, kicker, title, pill, meta)

    after_verdict = False
    i = 0
    while i < len(lines):
        raw = lines[i]
        s = raw.strip()
        if not s:
            i += 1
            continue
        if s.startswith("|") and i + 1 < len(lines) and TABLE_SEP.match(lines[i + 1].strip()):
            header = _cells(s)
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_cells(lines[i]))
                i += 1
            _data_table(doc, header, rows)
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            # The verdict gets its own labelled panel instead of a heading.
            after_verdict = _plain(m.group(2)).lower() == "verdict"
            if not after_verdict:
                _heading(doc, m.group(2), max(2, len(m.group(1))))
            i += 1
            continue
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
            rule = doc.add_paragraph()
            _para_border(rule, "bottom", RULE, 6)
            _spacing(rule, before=120, after=160, line=120)
            i += 1
            continue
        m = re.match(r"^(\s*)[-*+]\s+(.*)$", raw)
        if m:
            _bullet(doc, m.group(2), min(2, len(m.group(1)) // 2))
            i += 1
            continue
        m = re.match(r"^(\s*)(\d+(?:\.\d+)*|[a-z]|[ivx]+)[.)]\s+(.*)$", raw)
        if m:
            _clause(doc, m.group(2), m.group(3), min(2, len(m.group(1)) // 2))
            i += 1
            continue
        if s.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            _panel(doc, None, " ".join(buf))
            continue
        buf = [s]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||[-*+]\s|\d+[.)]\s|>)", lines[i].strip()):
            buf.append(lines[i].strip())
            i += 1
        text = " ".join(buf)
        v = re.match(r"^(GREEN|AMBER|RED)\b", _plain(text))
        if v or after_verdict:
            word = v.group(1) if v else ""
            fill, border, lab = {"RED": (MAGENTA_TINT, MAGENTA, MAGENTA), "AMBER": (CREAM, NAVY, NAVY),
                                 "GREEN": (PANEL, NAVY, NAVY)}.get(word, (CREAM, None, SLATE))
            _panel(doc, "Verdict", text, fill, border, lab)
            after_verdict = False
            continue
        _body(doc, text)

    _footer(sec, footer_text)
    doc.core_properties.title = title
    doc.core_properties.author = company or "VPS Lakeshore"
    out = io.BytesIO()
    doc.save(out)
    return _embed_fonts(out.getvalue())
