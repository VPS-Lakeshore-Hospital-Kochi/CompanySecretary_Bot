import io

import docx

from cs_assistant import export, files

SAMPLE = """# Non-Disclosure Agreement

This Agreement is made on [DATE].

## 1. Definitions

1. **Confidential Information** means all information.
2. *Purpose* means the evaluation.

- bullet one
  - nested

| No. | Clause | Risk |
|---|---|---|
| 1 | 5.2 | High |
"""


def test_markdown_to_docx_round_trip():
    data = export.markdown_to_docx(SAMPLE, "NDA", "Lakeshore Hospital")
    d = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Non-Disclosure Agreement" in text
    assert "1. Confidential Information means all information." in text
    assert d.tables and d.tables[0].rows[1].cells[2].text == "High"
    assert d.sections[0].header.paragraphs[0].text == "Lakeshore Hospital"


def test_docx_upload_reads_text_and_tables():
    src = docx.Document()
    src.add_paragraph("Clause 1. Term of three years.")
    t = src.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Fee"
    t.rows[0].cells[1].text = "Rs 5,00,000"
    buf = io.BytesIO()
    src.save(buf)
    doc = files.read_upload("agreement.docx", buf.getvalue())
    assert doc["kind"] == "text"
    assert "Term of three years" in doc["data"] and "Fee | Rs 5,00,000" in doc["data"]
    blocks = files.to_content_blocks(doc)
    assert blocks[0]["type"] == "document" and blocks[0]["source"]["type"] == "text"


def test_rejects_old_word_and_unknown():
    import pytest
    with pytest.raises(files.UploadError, match="Save As"):
        files.read_upload("old.doc", b"xx")
    with pytest.raises(files.UploadError):
        files.read_upload("x.exe", b"xx")
